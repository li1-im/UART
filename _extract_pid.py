from pptx import Presentation
import io

path = r"C:\Users\zyfdx\Desktop\2027\PID培训.pptx"
prs = Presentation(path)

out = io.StringIO()

def walk(shapes, indent="  "):
    for sh in shapes:
        try:
            if sh.__class__.__name__ == "GroupShape":
                out.write(f"{indent}[GROUP {sh.name}]\n")
                walk(sh.shapes, indent + "  ")
                continue
        except Exception:
            pass
        if sh.has_text_frame and sh.text_frame.text.strip():
            out.write(f"{indent}<{sh.name}> {sh.text_frame.text}\n")
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                out.write(indent + " | ".join(c.text.replace("\n", " ") for c in r.cells) + "\n")
        if sh.shape_type == 13:
            out.write(f"{indent}[PICTURE {sh.name}]\n")

for i, slide in enumerate(prs.slides, 1):
    out.write(f"\n{'='*70}\nSLIDE {i}\n{'='*70}\n")
    walk(slide.shapes)
    if slide.has_notes_slide:
        t = slide.notes_slide.notes_text_frame.text.strip()
        if t:
            out.write(f"-- NOTES: {t}\n")

with open(r"C:\Users\zyfdx\Desktop\STM32\CAN\_pid_text.txt", "w", encoding="utf-8") as f:
    f.write(out.getvalue())
print("ok", len(out.getvalue()))
